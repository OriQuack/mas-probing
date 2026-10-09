"""LLM client checks, budgets, configs, tools sandbox and leakage guards."""

from __future__ import annotations

import json

import pytest

from fakes import ScriptedLLM, fake_web
from minpilot.config import CONFIG_DIR, load_pool, load_refiner
from minpilot.data.gaia import GaiaTask
from minpilot.llm.client import LLMClient, TransientError
from minpilot.llm.specs import get_spec
from minpilot.runtime.trace import BudgetExceeded, InfraError, Limits, Trace
from minpilot.tools.config import REPO_ROOT
from minpilot.tools.toolbox import Toolbox


def client(tmp_path, key, llm, limits=None):
    return LLMClient(get_spec(key), Trace(tmp_path / "tr", limits), transport=llm, sleep=lambda s: None)


# -- client ------------------------------------------------------------------------------------
def test_request_pins_provider_and_reasoning(tmp_path):
    llm = ScriptedLLM({"worker": ["hi"]})
    client(tmp_path, "luna", llm).chat([{"role": "user", "content": "x"}])
    b = llm.bodies[0]
    assert b["provider"] == {"only": ["openai"], "allow_fallbacks": False, "require_parameters": True,
                             "data_collection": "deny"}
    assert b["reasoning"] == {"effort": "none"} and b["max_tokens"] == 8192


def test_effort_high_key_sends_tools_with_reasoning(tmp_path):
    """M1/M4 (2026-10-09): every role runs Luna at effort high, workers with tools (through OpenRouter)."""
    llm = ScriptedLLM({"worker": ["hi"]})
    client(tmp_path, "luna-high", llm).chat([{"role": "user", "content": "x"}],
                                            tools=[{"type": "function", "function": {"name": "f"}}])
    b = llm.bodies[0]
    assert b["reasoning"] == {"effort": "high"} and b["tools"] and b["max_tokens"] == 32768


def test_reasoning_tokens_counted_and_missing_reasoning_flagged(tmp_path):
    from minpilot.harness import reasoning_warnings

    tr = Trace(tmp_path / "tr")
    c = LLMClient(get_spec("luna-high"), tr, transport=ScriptedLLM({"worker": ["a", "b", "c"]}), sleep=lambda s: None)
    with tr.tags(stage="execution"):
        for _ in range(3):
            c.chat([{"role": "user", "content": "x"}])
    assert tr.usage["execution"]["reasoning_calls"] == 3 and tr.usage["execution"]["reasoning_tokens"] == 0
    assert reasoning_warnings(tr.usage) == ["no reasoning tokens in stage 'execution' (3 effort>none calls)"]
    tr.usage["execution"]["reasoning_tokens"] = 12
    assert reasoning_warnings(tr.usage) == []


def test_wrong_provider_is_infra_error(tmp_path):
    with pytest.raises(InfraError, match="served by provider"):
        client(tmp_path, "luna", ScriptedLLM({"worker": ["x"]}, provider="Azure")).chat([{"role": "user", "content": "x"}])


def test_missing_cost_is_infra_error(tmp_path):
    llm = ScriptedLLM({"worker": ["x"]})
    orig = llm.__call__

    def no_cost(body, t):
        r = orig(body, t)
        r["usage"].pop("cost")
        return r

    with pytest.raises(InfraError, match="no cost"):
        client(tmp_path, "luna", no_cost).chat([{"role": "user", "content": "x"}])


def test_transient_errors_retried_and_recorded(tmp_path):
    llm = ScriptedLLM({"worker": [TransientError("HTTP 503"), "ok"]})
    c = client(tmp_path, "luna", llm)
    assert c.chat([{"role": "user", "content": "x"}]).content == "ok"
    recs = [json.loads(l) for l in (tmp_path / "tr" / "llm_calls.jsonl").read_text().splitlines()]
    assert [r["status"] for r in recs] == ["started", "error", "started", "ok"]
    assert c.trace.run_scope.llm_calls == 2


def test_cost_cap_refuses_before_sending(tmp_path):
    llm = ScriptedLLM({"worker": ["x"]})
    with pytest.raises(BudgetExceeded):
        client(tmp_path, "luna", llm, Limits(max_cost_usd=0.001)).chat([{"role": "user", "content": "x"}])
    assert llm.bodies == []  # the worst case (8192 output tokens) did not fit; nothing was sent


def test_reasoning_details_kept_for_gemini(tmp_path):
    def reply(body, t):
        return {"id": "g", "provider": "Google AI Studio", "model": body["model"],
                "choices": [{"message": {"role": "assistant", "content": "a", "reasoning_details": [{"x": 1}]}}],
                "usage": {"cost": 1e-6}}

    res = client(tmp_path, "gemini", reply).chat([{"role": "user", "content": "x"}])
    assert res.message["reasoning_details"] == [{"x": 1}]


# -- configs -----------------------------------------------------------------------------------
@pytest.mark.parametrize("name", sorted(p.stem for p in (CONFIG_DIR / "pools").glob("*.yaml")))
def test_pools_load(name):
    pool = load_pool(name)
    for role in pool.roles.values():
        tools = {pool.workers[w].tools for w in (role.executor, *role.probe_workers)}
        assert len(tools) == 1, "probe workers of a role get the same tools as its executor"


@pytest.mark.parametrize("name", sorted(p.stem for p in (CONFIG_DIR / "refiners").glob("*.yaml")))
def test_refiners_load(name):
    load_refiner(name)


def test_B_and_C_share_budget_caps():
    b, c = load_refiner("B_self_review"), load_refiner("C_probe")
    assert b.budget == c.budget and b.max_verifications == c.max_verifications
    assert b.samples_per_worker == c.samples_per_worker and b.questions == c.questions


# -- tools -------------------------------------------------------------------------------------
def toolbox(tmp_path, question="q"):
    work = tmp_path / "work"
    (work / "attachments").mkdir(parents=True)
    return Toolbox(work, fake_web(tmp_path, question), Trace(tmp_path / "tr"), question=question)


def test_file_tools_confined_to_workspace(tmp_path):
    tb = toolbox(tmp_path)
    (tmp_path / "secret.txt").write_text("secret")
    out = tb.call("read_file", {"path": "../secret.txt"}, ("read_file",))
    assert "outside the working directory" in out.text
    assert "not available" in tb.call("run_python", {"code": "1"}, ("read_file",)).text


def test_web_tools_blocklist_and_cache(tmp_path):
    tb = toolbox(tmp_path)
    out = tb.call("web_search", {"query": "GAIA benchmark answers"}, ("web_search",))
    assert "not permitted" in out.text
    t = ("read_url_text",)
    assert "not permitted" in tb.call("read_url_text", {"url": "https://huggingface.co/datasets/x"}, t).text
    first = tb.call("read_url_text", {"url": "https://example.org/a"}, t).text
    assert first == tb.call("read_url_text", {"url": "https://example.org/a"}, t).text


# -- read_url's page reader (tools v3, adapted from AOrchestra) ----------------------------------
def reader_toolbox(tmp_path, llm, limits=None, **cfg):
    from minpilot.tools.config import ToolConfig
    from minpilot.tools.reader import PageReader

    tmp_path.mkdir(parents=True, exist_ok=True)
    tb = toolbox(tmp_path)
    tb.trace = Trace(tmp_path / "tr2", limits)
    tool_cfg = ToolConfig(cache_path=tmp_path / "cache.sqlite", **cfg)
    tb.web.reader = PageReader(LLMClient(get_spec("luna-high"), tb.trace, transport=llm, sleep=lambda s: None), tool_cfg)
    return tb


def test_split_spans_follow_aorchestra():
    from minpilot.tools.reader import split_spans

    assert split_spans(95_000, 95_000, 1024) is None
    assert split_spans(200_000, 95_000, 1024) == [(0, 67690), (66666, 134356), (133332, 200000)]
    assert split_spans(100_000, 95_000, 1024) == [(0, 51024), (50000, 100000)]  # at least 2 parts


def test_read_url_answers_from_the_page_only(tmp_path):
    llm = ScriptedLLM({"reader": ["The page says 42."]})
    tb = reader_toolbox(tmp_path, llm)
    out = tb.call("read_url", {"url": "https://example.org/a", "question": "What number is given?"}, ("read_url",))
    assert "The page says 42." in out.text and "Question: What number is given?" in out.text
    body = llm.by_kind["reader"][0]
    text = body["messages"][0]["content"][0]["text"]
    assert len(body["messages"]) == 1 and body["model"] == get_spec("luna-high").model
    assert body["reasoning"] == {"effort": "high"}
    assert text.startswith("Please read the source content") and "Text of https://example.org/a" in text
    assert text.endswith("What number is given?") and "unique orders" not in text  # never the task
    rec = [json.loads(l) for l in (tmp_path / "tr2" / "tool_calls.jsonl").read_text().splitlines()][-1]
    assert rec["status"] == "ok" and rec["reader"]["calls"] == 1 and rec["reader"]["n_parts"] == 1
    assert rec["reader"]["usd"] > 0 and rec["question"] == "What number is given?"
    calls = [json.loads(l) for l in (tmp_path / "tr2" / "llm_calls.jsonl").read_text().splitlines()]
    assert all(c["component"] == "reader" and c["tool_call"] == rec["call_no"] for c in calls)
    assert tb.trace.run_scope.llm_usd > 0  # the reader's cost counts toward the run


def test_read_url_splits_long_pages(tmp_path):
    llm = ScriptedLLM({"reader": lambda body: "part answer"})
    tb = reader_toolbox(tmp_path, llm, reader_part_tokens=400, reader_overlap_tokens=10)
    out = tb.call("read_url", {"url": "https://example.org/a", "question": "q?"}, ("read_url",))
    n = len(llm.by_kind["reader"])
    assert n >= 3 and "Since the content is too long" in out.text and f"result part {n} ---" in out.text
    assert all("--- begin of source content ---" in b["messages"][0]["content"][0]["text"]
               for b in llm.by_kind["reader"])


def test_read_url_reader_failure_and_budget(tmp_path):
    tb = reader_toolbox(tmp_path / "a", ScriptedLLM({"reader": [""]}))
    out = tb.call("read_url", {"url": "https://example.org/a", "question": "q?"}, ("read_url",))
    assert out.text.startswith("Error: the page reader failed")
    tb = reader_toolbox(tmp_path / "b", ScriptedLLM({"reader": ["x"]}), limits=Limits(max_llm_calls=0))
    with pytest.raises(BudgetExceeded):
        tb.call("read_url", {"url": "https://example.org/b", "question": "q?"}, ("read_url",))
    rec = [json.loads(l) for l in (tmp_path / "b" / "tr2" / "tool_calls.jsonl").read_text().splitlines()][-1]
    assert rec["status"] == "error" and rec["error"].startswith("BudgetExceeded")


def test_find_in_url(tmp_path):
    tb = toolbox(tmp_path)
    out = tb.call("find_in_url", {"url": "https://example.org/a", "text": "TEXT   of"}, ("find_in_url",)).text
    assert "50 occurrence(s)" in out and "[1] page 1, character 0" in out and "Showing the first 20" in out
    assert "0 occurrence(s)" in tb.call("find_in_url", {"url": "https://example.org/a", "text": "zebra"},
                                        ("find_in_url",)).text


def test_view_image_returns_image(tmp_path):
    from PIL import Image

    tb = toolbox(tmp_path)
    Image.new("RGB", (40, 20), "red").save(tb.work_dir / "attachments" / "x.png")
    out = tb.call("view_image", {"path": "attachments/x.png"}, ("view_image",))
    assert out.images and out.images[0].startswith("data:image/png;base64,") and "40x20" in out.text


@pytest.mark.landlock
def test_sandbox_blocks_repo_network_and_processes(tmp_path):
    tb = toolbox(tmp_path)
    run = lambda code: tb.call("run_python", {"code": code}, ("run_python",)).text  # noqa: E731
    assert "Permission denied" in run(f"print(open({str(REPO_ROOT / 'pyproject.toml')!r}).read())")
    assert "Network access is not permitted" in run("import urllib.request; urllib.request.urlopen('http://example.org')")
    assert "Starting processes is not permitted" in run("import os; os.system('ls')")
    assert "hello" in run("open('f.txt','w').write('hello'); print(open('f.txt').read())")
    assert str(tb.work_dir) not in run("import os; print(os.getcwd())")


@pytest.mark.landlock
@pytest.mark.gaia_data
def test_sandbox_cannot_read_gaia_answers(tmp_path):
    from minpilot.data.gaia import GAIA_ROOT

    tb = toolbox(tmp_path)
    p = GAIA_ROOT / "validation" / "metadata.parquet"
    out = tb.call("run_python", {"code": f"import pandas as pd; print(pd.read_parquet({str(p)!r}).shape)"},
                  ("run_python",)).text
    assert "Permission denied" in out or "Error" in out


@pytest.mark.gaia_data
def test_tasks_carry_no_ground_truth():
    from minpilot.data.gaia import load_tasks

    t = load_tasks()[0]
    assert set(GaiaTask.__dataclass_fields__) == {"task_id", "level", "question", "file_name", "file_path"}
    assert t.question


def test_error_outputs_are_recorded_as_errors(tmp_path):
    tb = toolbox(tmp_path)
    tb.call("read_file", {"path": "missing.txt"}, ("read_file",))
    recs = [json.loads(l) for l in (tmp_path / "tr" / "tool_calls.jsonl").read_text().splitlines()]
    done = [r for r in recs if r["status"] != "started"]
    assert done[-1]["status"] == "error" and done[-1]["error"]


def test_blocklist_v4_query_echo_and_qa_aggregator_pages():
    from minpilot.tools import blocklist

    assert blocklist.url_block_reason("https://nuggetpedia.com/nugget/nw-1") == "qa_aggregator"
    assert blocklist.url_block_reason("https://www.instagram.com/popular/some-question-words/") == "query_echo_page"
    assert blocklist.url_block_reason("https://www.instagram.com/natgeo/") is None


def test_docx_tables_stay_in_document_order():
    import io

    import docx

    from minpilot.tools.documents import docx_to_text

    d = docx.Document()
    d.add_paragraph("Heading A")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "x", "y"
    d.add_paragraph("After table")
    buf = io.BytesIO()
    d.save(buf)
    lines = docx_to_text(buf.getvalue()).splitlines()
    assert lines == ["Heading A", "[table 1]", "x | y", "After table"]


def test_blocklist_v5_rules():
    from minpilot.tools import blocklist

    q = ("In the endnote found in the second-to-last paragraph of page 11 of the book with the doi 10.2307/j.ctv9b2xdv, "
         "what date in November was the Wikipedia article accessed?")
    hyph = ("... In the end- note found in the second-to-last paragraph of page 11 of the book with the doi "
            "10.2307/j.ctv9b2xdv ...")
    assert blocklist.content_block_reason(hyph, q) == "quotes_task_question"
    trace = "# --- Sub-task 1: WebSearch for the endnote Wikipedia article accessed November date"
    assert blocklist.content_block_reason(trace, q) == "agent_trace_on_task"
    assert blocklist.content_block_reason("Sub-task 1: buy groceries for the weekend", q) is None


def test_blocked_url_stays_blocked_for_the_run(tmp_path):
    from minpilot.tools.backends import SearchHit

    q = "What is the airspeed velocity of an unladen swallow according to the castle guard in the film?"
    web = fake_web(tmp_path, q)
    leak = SearchHit(title="Agent trace", url="https://example.org/paper",
                     snippet="multi-agent trace: what is the airspeed velocity of an unladen swallow according to")
    benign = SearchHit(title="Paper", url="https://example.org/paper", snippet="A paper about birds.")
    web.search_backends[0].results = {"first": [leak], "second": [benign]}
    out1, rec1 = web.web_search("first")
    out2, rec2 = web.web_search("second")
    assert rec1["blocked_reasons"] and rec2["blocked_reasons"] == ["blocked_earlier_in_run"]
    assert "example.org/paper" not in out2
    assert web.state_dict()["blocked_urls"] == ["https://example.org/paper"]


def test_search_fallback_library_importable():
    """The DDG fallback imports its library lazily; a missing package only shows up when Serper fails."""
    try:
        from ddgs import DDGS  # noqa: F401
    except ImportError:
        from duckduckgo_search import DDGS  # noqa: F401


def test_read_url_date_ignored_for_snapshot_urls(tmp_path):
    class NoWayback:
        def closest(self, url, date):
            raise AssertionError("a snapshot URL must not be looked up again")

    web = fake_web(tmp_path)
    web.wayback = NoWayback()
    url = "https://web.archive.org/web/20210105000000/https://en.wikipedia.org/wiki/Greenland"
    text, rec = web.read_url_text(url, date="20210101")
    assert "Text of" in text and not rec.get("error")


def _score_module():
    import importlib.util

    spec = importlib.util.spec_from_file_location("score_runs", REPO_ROOT / "scripts" / "score_runs.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_scoring_groups_reps_by_condition_and_pairs_checkpoints():
    import pandas as pd

    sr = _score_module()
    base = {"mode": "restore", "pool": "role_routing", "override_source": None}
    rows = [
        {**base, "task_id": "t1", "condition": "C_probe", "label": "C_r0", "checkpoint": "ck1", "correct": True},
        {**base, "task_id": "t1", "condition": "C_probe", "label": "C_r1", "checkpoint": "ck1", "correct": False},
        {**base, "task_id": "t1", "condition": "A_none", "label": "A_r0", "checkpoint": "ck1", "correct": False},
        {**base, "task_id": "t2", "condition": "C_probe", "label": "C_r0", "checkpoint": "ck2", "correct": True},
        {**base, "task_id": "t2", "condition": "A_none", "label": "A_r0", "checkpoint": "ck2", "correct": True},
        {**base, "task_id": "t2", "condition": "A_none", "label": "post", "checkpoint": "ck2", "correct": True,
         "override_source": "post_hoc"},
    ]
    c_rows = [r for r in rows if sr.arm_of(r) == "restore|role_routing|C_probe|override=none"]
    assert sr.per_task(c_rows) == {"t1": 0.5, "t2": 1.0}          # C_r0 and C_r1 are one condition
    res = sr.compare(pd.DataFrame(rows), "A_none", "C_probe")["role_routing"]
    assert res["tasks"] == 2 and res["x"] == 0.5 and res["y"] == 0.75 and res["diff"] == 0.25
    assert res["ci95"][0] <= 0.25 <= res["ci95"][1]


def test_paid_tool_calls_count_in_usd_and_budget(tmp_path):
    from minpilot.runtime.costs import TOOL_USD_PER_CREDIT

    tr = Trace(tmp_path / "tr", Limits(max_cost_usd=1.0))
    tr.tool({"tool": "web_search", "status": "ok", "backend": "serper", "credits": {"serper": 1}, "cache_hit": False})
    tr.tool({"tool": "web_search", "status": "ok", "backend": "serper", "credits": {}, "cache_hit": True})
    c = tr.run_scope.counters()
    price = TOOL_USD_PER_CREDIT["serper"]
    assert c["tool_usd"] == price and c["tool_usd_cold"] == 2 * price and c["cost_usd"] == price
    tr0 = Trace(tmp_path / "tr0", Limits(max_cost_usd=0.0))
    with pytest.raises(BudgetExceeded):
        tr0.before_tool()


def test_unbilled_llm_attempts_are_counted(tmp_path):
    tr = Trace(tmp_path / "tr", Limits())
    tr.settle_llm(0.01, 0.0, unknown=True)
    assert tr.run_scope.counters()["unknown_cost_calls"] == 1


def test_served_model_must_match_the_request(tmp_path):
    from minpilot.llm.client import same_model

    assert same_model("openai/gpt-6-luna", "openai/gpt-6-luna-20260922")
    assert same_model("openai/gpt-4o-2024-08-06", "openai/gpt-4o-2024-08-06")
    assert not same_model("openai/gpt-4o", "openai/gpt-6-luna-20260922")
    tr = Trace(tmp_path / "tr", Limits())
    client = LLMClient(get_spec("luna"), tr, transport=ScriptedLLM({"worker": ["hi"]}, served_model="openai/gpt-4o"),
                       sleep=lambda s: None)
    with pytest.raises(InfraError, match="served model"):
        client.chat([{"role": "user", "content": "x"}])


def test_scratch_web_state_does_not_change_the_main_runs(tmp_path):
    web = fake_web(tmp_path)
    scratch = web.fork()
    scratch.web_search("probe query")
    scratch.blocked_urls.add("https://example.org/x")
    assert scratch.pinned_search == "serper" and web.pinned_search is None
    assert web.blocked_urls == set() and scratch.cache is web.cache


def test_online_images_are_cached(tmp_path, monkeypatch):
    import io

    from PIL import Image

    from minpilot.tools import backends

    calls = []

    def download(self, url):
        buf = io.BytesIO()
        Image.new("RGB", (10 + len(calls), 10), "red").save(buf, format="PNG")
        calls.append(url)
        return buf.getvalue(), "image/png", url

    monkeypatch.setattr(backends.DirectFetchBackend, "download", download)
    tb = toolbox(tmp_path)
    a = tb.call("view_image", {"path": "https://example.org/i.png"}, ("view_image",))
    b = tb.call("view_image", {"path": "https://example.org/i.png"}, ("view_image",))
    assert len(calls) == 1 and a.images == b.images
