"""Model-free conversion of fetched bytes and local files to text.

No model inside any reader: long outputs are paged by the caller, never summarized. Copied unchanged from the
previous pilot.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import html2text

AV_EXTS = {"mp3", "wav", "m4a", "flac", "ogg", "mp4", "mov", "avi", "mkv", "webm"}
IMAGE_EXTS = {"png", "jpg", "jpeg", "gif", "bmp", "webp", "tif", "tiff"}
TEXT_EXTS = {"txt", "md", "py", "json", "jsonl", "jsonld", "xml", "csv", "tsv", "pdb", "html", "htm", "yaml",
             "yml", "js", "css", "tex", "log", "ini", "cfg", "toml"}

_CONTENT_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel": "xls",
    "application/zip": "zip",
    "text/csv": "csv",
    "application/json": "json",
    "text/plain": "txt",
    "application/xml": "xml",
    "text/xml": "xml",
}


class UnsupportedDocument(Exception):
    pass


def ext_from_content_type(content_type: str, url: str = "") -> str:
    ct = content_type.split(";")[0].strip().lower()
    if ct in ("text/html", "application/xhtml+xml"):
        return "html"
    if ct in _CONTENT_TYPES:
        return _CONTENT_TYPES[ct]
    if ct.startswith("image/"):
        return ct.split("/")[1]
    suffix = Path(url.split("?")[0]).suffix.lower().lstrip(".")
    return suffix or "html"


def html_to_text(html: str) -> str:
    h = html2text.HTML2Text()
    h.body_width = 0  # no hard wrapping
    h.ignore_images = True
    h.ignore_links = False
    h.ignore_tables = False
    return h.handle(html).strip()


def html_title(html: str) -> str:
    from bs4 import BeautifulSoup

    title = BeautifulSoup(html, "lxml").title
    return title.get_text(strip=True) if title else ""


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("latin-1")


def pdf_to_text(data: bytes) -> str:
    import fitz

    with fitz.open(stream=data, filetype="pdf") as doc:
        return "\n\n".join(f"[page {i + 1}]\n{page.get_text().strip()}" for i, page in enumerate(doc))


def docx_to_text(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t_idx, table in enumerate(d.tables, 1):
        parts.append(f"[table {t_idx}]")
        parts += [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
    return "\n".join(parts)


def pptx_to_text(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"[slide {i}]")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                out += [" | ".join(c.text.strip() for c in row.cells) for row in shape.table.rows]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append(f"(notes) {slide.notes_slide.notes_text_frame.text.strip()}")
    return "\n".join(out)


def spreadsheet_to_text(data: bytes, ext: str) -> str:
    """All sheets as CSV, plus non-default cell fill colours (some GAIA tasks are about coloured cells)."""
    import pandas as pd

    if ext == "csv":
        return _decode(data)
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None)
    out = []
    for name, df in sheets.items():
        out.append(f"[sheet {name!r}: {df.shape[0]} rows x {df.shape[1]} cols, no header inferred]")
        out.append(df.to_csv(index=False, header=False).strip())
    if ext == "xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data))
        for ws in wb.worksheets:
            fills = []
            for row in ws.iter_rows():
                for cell in row:
                    fill = cell.fill
                    rgb = getattr(fill.fgColor, "rgb", None) if fill and fill.fill_type == "solid" else None
                    if isinstance(rgb, str) and rgb not in ("00000000", "FFFFFFFF"):
                        fills.append(f"{cell.coordinate}={rgb}({cell.value!r})")
            if fills:
                out.append(f"[cell fill colours in sheet {ws.title!r} (ARGB)]")
                out.append(", ".join(fills))
    return "\n".join(out)


def bytes_to_text(data: bytes, ext: str) -> str:
    ext = ext.lower()
    if ext in ("html", "htm"):
        return html_to_text(_decode(data))
    if ext == "pdf":
        return pdf_to_text(data)
    if ext == "docx":
        return docx_to_text(data)
    if ext == "pptx":
        return pptx_to_text(data)
    if ext in ("xlsx", "xls", "csv"):
        return spreadsheet_to_text(data, ext)
    if ext in TEXT_EXTS:
        return _decode(data)
    if ext in IMAGE_EXTS:
        raise UnsupportedDocument(f"'{ext}' is an image; text extraction is not available for images.")
    if ext in AV_EXTS:
        raise UnsupportedDocument(f"'{ext}' audio/video files are not supported.")
    raise UnsupportedDocument(f"Unsupported file type '{ext}'.")


def safe_dir(root: Path, rel: str | Path) -> Path:
    """`root/rel` as a real directory inside `root`: refuses a path that is (or passes through) a symlink
    leading outside. Tools write here from the runner process, which the code sandbox does not cover."""
    root = Path(root).resolve()
    target = root / rel
    if target.is_symlink() or not target.resolve().is_relative_to(root):
        raise PermissionError(f"{rel} resolves outside the working directory")
    target.mkdir(parents=True, exist_ok=True)
    if not target.resolve().is_relative_to(root):
        raise PermissionError(f"{rel} resolves outside the working directory")
    return target


def extract_zip(data: bytes, dest: Path, root: Path | None = None) -> list[Path]:
    """Extract into dest and return the extracted files. dest must resolve inside `root` (the work dir;
    default: dest itself), every entry must resolve inside dest, and no existing symlink is written
    through, so neither `../` entries nor a symlinked dest or sub-directory can make it write elsewhere."""
    root = Path(root or dest).resolve()
    dest = Path(dest)
    if dest.is_symlink() or not dest.resolve().is_relative_to(root):
        raise PermissionError(f"extraction target {dest} resolves outside the working directory")
    dest.mkdir(parents=True, exist_ok=True)
    dest = dest.resolve()
    files = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            target = dest / info.filename
            # Entries must stay inside dest (zip-slip: "../x"), and resolve() follows existing symlinks along
            # the path, so a symlinked sub-directory pointing elsewhere is skipped too. dest is inside root.
            if not target.resolve().is_relative_to(dest) or any(
                    p.is_symlink() for p in [target, *target.parents] if p.is_relative_to(dest) and p != dest):
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            if not target.parent.resolve().is_relative_to(dest):
                continue
            target.write_bytes(zf.read(info))
            files.append(target)
    return files
